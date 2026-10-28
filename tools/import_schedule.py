#!/usr/bin/env python3
"""Read an old music schedule spreadsheet (.xlsx) into a plain table, so the planner knows what was done.

    python tools/import_schedule.py parish/schedule/Music_Schedule.xlsx

Writes parish/schedule/schedule.csv with one row per item: date, occasion, organist, slot, item.

It reads the layout tools/spreadsheet.py writes: on every tab, a column per service with its date in the
first row, and a row per piece named in the first column (your names mapped by spreadsheet_rows: in
parish.txt), with rows called Proper (the day) and Organist. A spreadsheet laid out otherwise needs this
script adapted; it's short.
"""
import csv, datetime as dt, os, sys

TOOLS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, TOOLS)
import paths, spreadsheet   # noqa: E402


def main():
    import openpyxl
    src = sys.argv[1] if len(sys.argv) > 1 else spreadsheet.path()
    wb = openpyxl.load_workbook(src, data_only=True)
    rows = []
    for ws in wb.worksheets:
        labels = {r: (ws.cell(row=r, column=1).value or '').strip() for r in range(1, ws.max_row + 1)}
        labels = {r: l for r, l in labels.items() if l}
        for c in range(2, ws.max_column + 1):
            date = ws.cell(row=1, column=c).value
            if not isinstance(date, dt.datetime):
                continue
            col = {labels[r]: ws.cell(row=r, column=c).value for r in labels if r > 1}
            col = {spreadsheet.canon(k): v for k, v in col.items()}
            occasion = str(col.get('Proper') or '').strip()
            organist = str(col.get('Organist') or '').strip()
            for label, val in col.items():
                if label in ('Proper', 'Organist') or val in (None, ''):
                    continue
                slot = label
                rows.append({'date': date.date().isoformat(), 'occasion': occasion, 'organist': organist,
                             'slot': slot, 'item': str(val).strip()})
    rows.sort(key=lambda r: r['date'])
    out = os.path.join(paths.HOME, 'schedule', 'schedule.csv')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, 'w', encoding='utf-8', newline='') as fh:
        w = csv.DictWriter(fh, ['date', 'occasion', 'organist', 'slot', 'item'])
        w.writeheader()
        w.writerows(rows)
    print(f'{len(rows)} items, {len({r["date"] for r in rows})} services, {rows[0]["date"]} to {rows[-1]["date"]}')


if __name__ == '__main__':
    main()
