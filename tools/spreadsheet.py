#!/usr/bin/env python3
"""Write a Sunday's music into the music schedule spreadsheet (the .xlsx kept since 2022).

    python tools/spreadsheet.py parish/sundays/2026-11-08.txt          # show what would change
    python tools/spreadsheet.py parish/sundays/2026-11-08.txt --write  # write it

The spreadsheet has one sheet per year, one column per service (the date in row 1) and one row per
piece (Prelude, Opening, Gloria, Psalm, …). Which row each piece of a service goes in comes from the
[brackets] in parish/services.txt. Only cells the Sunday file has something for are touched, and
before writing, a copy of the spreadsheet is kept in schedule/backups/.

Where the spreadsheet is: `spreadsheet:` in parish.txt. `google` means a Google Sheet, reached through
the small script in tools/google_sheet.gs (connect it once from the planner); anything else is the path
of an .xlsx file (default: schedule/SP_Music_Schedule.xlsx in the parish folder; close it in Excel first).
"""
import copy, datetime as dt, glob, os, re, shutil, sys

TOOLS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, TOOLS)
import paths, service   # noqa: E402

# old row labels -> the slot names used in services.txt's [brackets] (same as tools/import_schedule.py)
ROW_ALIASES = {'opening hymn': 'Opening', 'gradual hymn': 'Gradual', 'communion hymn': 'Communion',
               'closing hymn': 'Closing', 'anthem': 'Offertory'}
BOOK_IN_SHEET = {'LEVAS': 'LEVAS2'}          # how the spreadsheet writes the hymnals
HYMNARY = {'EH1982': 'EH1982', 'LEVAS': 'LEVS1993', 'WLP': 'WLP1997'}    # its hymn cells link to hymnary.org


def link_for(value):
    """The hymnary.org page of the first hymn in a cell, as the spreadsheet's hymn cells have."""
    m = re.search(r'\b(S?\d{1,3}[a-z]?) \((EH1982|LEVAS2?|WLP)\)', value or '')
    if not m:
        return None
    book = HYMNARY['LEVAS' if m.group(2).startswith('LEVAS') else m.group(2)]
    return f'https://hymnary.org/hymn/{book}/{m.group(1).lower()}'


def set_cell(cell, value, fonts):
    """Write a value; hymn cells link to hymnary.org and look like links, the rest look plain."""
    cell.hyperlink = None
    link = link_for(value)
    if link:
        cell.hyperlink = link
    cell.value = value
    f = fonts.get('link' if link else 'plain')
    if f is not None:
        cell.font = copy.copy(f)


def fonts_of(ws, rows):
    """The look of a linked cell and of a plain one in this sheet, to copy."""
    out = {}
    for row in ws.iter_rows(min_row=2, max_row=max(rows.values())):
        for c in row:
            if c.column == 1 or c.value in (None, '') or type(c).__name__ == 'MergedCell':
                continue
            k = 'link' if c.hyperlink else 'plain'
            if k not in out and (k == 'link' or not c.font.underline):
                out[k] = c.font
        if len(out) == 2:
            break
    return out
KEEP_BACKUPS = 20


def path():
    import lectionary
    p = lectionary.PARISH.get('spreadsheet', '').strip() or os.path.join('schedule', 'SP_Music_Schedule.xlsx')
    p = os.path.expanduser(p)
    return p if os.path.isabs(p) else os.path.join(paths.HOME, p)


def sheet_hymns(hs):
    return ', '.join(f"{h['number']} ({BOOK_IN_SHEET.get(h['book'], h['book'])})" for h in hs)


def cell_value(s, v):
    """A piece as the spreadsheet writes it: '390 (EH1982)', 'S278 (EH1982)', or the text as typed."""
    hs = service.hymns(v) if s['kind'] in ('hymn', 'anthem', 'service') else []
    if not hs:
        return v
    out = sheet_hymns(hs)
    return ('Kyrie ' + out) if s['key'] == 'kyrie' else out


def psalm_value(spec, chants):
    names = [n.replace('tone:', 'Tone ') for n, _ in chants if n not in ('???', '?')]
    return spec + (f" ({', '.join(names)})" if names else '')


def wanted(sunday, current):
    """{row slot name: new value} for this Sunday. current = {slot name: what the cell holds now}."""
    order = service.order_of(sunday)
    sl = service.slots(order)
    given = dict(service.music(sunday))
    cells = {}                        # slot name -> [(slot, value)]
    for s in sl:
        v = given.get(s['key'], '').strip()
        if not v or not s['columns']:
            continue
        col = next((c for c in s['columns'] if (current.get(c) or '').strip()), s['columns'][0]) if len(s['columns']) > 1 else s['columns'][0]
        cells.setdefault(col, []).append((s, v))
    out = {}
    for col, pieces in cells.items():
        if len(pieces) == 1:
            out[col] = cell_value(*pieces[0])
            continue
        vals = [cell_value(s, v) for s, v in pieces]
        m = [re.match(r'^(Magnificat|Nunc dimittis)\b(.*)$', v) for v in vals]
        if len(vals) == 2 and all(m) and m[0].group(1) != m[1].group(1) and m[0].group(2) == m[1].group(2):
            out[col] = 'Magnificat & Nunc dimittis' + m[0].group(2)        # one setting of both canticles
        else:
            out[col] = ', '.join(f"{s['label']}: {val}" for (s, _), val in zip(pieces, vals))
    # "also" lines that say where they belong ("Communion hymn: 204") join that row
    for k, v in service.music(sunday):
        if k != 'also' or ':' not in v:
            continue
        lab, val = (x.strip() for x in v.split(':', 1))
        col = next((c for c in set(ROW_ALIASES.values()) | {'Prelude', 'Postlude', 'Gloria', 'Psalm'}
                    if lab.lower().split()[0] == c.lower().split()[0]), None)
        if col:
            hs = service.hymns(val)
            item = f"{lab}: {sheet_hymns(hs) if hs else val}"
            out[col] = (out.get(col) or (current.get(col) or '').strip())
            out[col] = f"{out[col]}; {item}" if out[col] and item not in out[col] else (out[col] or item)
    if sunday['items']:
        out['Psalm'] = '; '.join(psalm_value(it['spec'], it['chants']) for it in sunday['items'])
    org = given.get('organist', '').strip()
    if org:
        out['Organist'] = org
    return out


def find_column(ws, date):
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if isinstance(v, dt.datetime) and v.date() == date:
            return c
    return None


def rows_of(ws):
    out = {}
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(row=r, column=1).value
        if isinstance(lab, str) and lab.strip():
            out.setdefault(ROW_ALIASES.get(lab.strip().lower(), lab.strip()), r)
    return out


def date_columns(ws):
    return [c for c in range(2, ws.max_column + 1) if isinstance(ws.cell(row=1, column=c).value, dt.datetime)]


def insert_column(ws, date):
    """Make room for a service the sheet doesn't have yet, in date order, laid out like its neighbours
    (each service takes two columns, the date merged across them)."""
    from openpyxl.utils import get_column_letter
    cols = date_columns(ws)
    stamp = dt.datetime(date.year, date.month, date.day)
    merged_at = lambda c: any(m.min_row == 1 and m.min_col == c for m in ws.merged_cells.ranges)
    if not cols:                                     # a new year's sheet: the first service
        ws.cell(row=1, column=2).value = stamp
        return 2
    width = (cols[1] - cols[0]) if len(cols) > 1 else 2
    later = [c for c in cols if ws.cell(row=1, column=c).value.date() > date]
    if not later and merged_at(cols[-1] + width) and ws.cell(row=1, column=cols[-1] + width).value is None:
        ws.cell(row=1, column=cols[-1] + width).value = stamp     # an empty column laid out already
        ws.cell(row=1, column=cols[-1] + width).number_format = ws.cell(row=1, column=cols[-1]).number_format
        return cols[-1] + width
    at = later[0] if later else cols[-1] + width
    model = (later[0] if later else cols[-1])
    # move everything from `at` rightwards by hand (openpyxl's insert_cols leaves merges and widths behind)
    merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    last = ws.max_column
    widths = {c: ws.column_dimensions[get_column_letter(c)].width for c in range(1, last + 1)}
    if later:
        for r in range(1, ws.max_row + 1):
            for c in range(last, at - 1, -1):
                src, dst = ws.cell(row=r, column=c), ws.cell(row=r, column=c + width)
                dst.hyperlink = None
                if src.hyperlink:
                    dst.hyperlink = src.hyperlink.target
                dst.value = src.value
                dst._style = copy.copy(src._style)
                dst.number_format = src.number_format
        for c in range(last, at - 1, -1):
            if widths.get(c):
                ws.column_dimensions[get_column_letter(c + width)].width = widths[c]
    msrc = model + width if later else model           # the neighbour, after the move
    for r in range(1, ws.max_row + 1):
        for k in range(width):
            src, dst = ws.cell(row=r, column=msrc + k), ws.cell(row=r, column=at + k)
            dst.hyperlink = None
            dst.value = None
            dst._style = copy.copy(src._style)
    for k in range(width):
        if widths.get(model + k):
            ws.column_dimensions[get_column_letter(at + k)].width = widths[model + k]
    for (r0, c0, r1, c1) in merges:
        if later and c0 >= at:
            c0, c1 = c0 + width, c1 + width
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
        if c0 == msrc and c1 == msrc + width - 1:      # lay the new service out the same way
            ws.merge_cells(start_row=r0, start_column=at, end_row=r1, end_column=at + width - 1)
    ws.cell(row=1, column=at).value = dt.datetime(date.year, date.month, date.day)
    ws.cell(row=1, column=at).number_format = ws.cell(row=1, column=msrc).number_format
    return at


def new_year_sheet(wb, year):
    """A sheet for a new year, with the row labels and look of the latest one and no services yet."""
    last = max((ws for ws in wb.worksheets if ws.title.isdigit()), key=lambda w: int(w.title))
    ws = wb.copy_worksheet(last)
    ws.title = str(year)
    for c in range(2, ws.max_column + 1):                  # keep the layout, empty every service
        for r in range(1, ws.max_row + 1):
            cell = ws.cell(row=r, column=c)
            if type(cell).__name__ != 'MergedCell':
                cell.hyperlink = None
                cell.value = None
    wb.move_sheet(ws, offset=-(len(wb.worksheets) - 1))   # newest year first, as now
    return ws


# ------------------------------------------------------------------ where the spreadsheet is
SHEET_SECRET = 'psalter-sheet'


def canon(label):
    return ROW_ALIASES.get(label.strip().lower(), label.strip())


class XlsxSheet:
    """The spreadsheet as an .xlsx file on this computer."""
    def __init__(self, p):
        self.path, self.name = p, os.path.basename(p)
        if not os.path.exists(p):
            raise RuntimeError(f'no spreadsheet at {p} (set spreadsheet: in parish.txt)')

    def read(self, date):
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        ws = wb[str(date.year)] if str(date.year) in wb.sheetnames else None
        model = ws or max((w for w in wb.worksheets if w.title.isdigit()), key=lambda w: int(w.title))
        rows = rows_of(model)
        col = find_column(ws, date) if ws else None
        cells = {}
        if col:
            for name, r in rows.items():
                v = ws.cell(row=r, column=col).value
                cells[name] = '' if v is None else (v.strftime('%m/%d/%Y') if isinstance(v, dt.datetime) else str(v))
        labels = {canon(ws_label): ws_label for ws_label in
                  [model.cell(row=r, column=1).value.strip() for r in rows.values()]}
        return {'sheet': ws is not None, 'column': col is not None, 'labels': labels, 'cells': cells}

    def write(self, date, cells):
        import openpyxl
        p = self.path
        bdir = os.path.join(os.path.dirname(p), 'backups')
        os.makedirs(bdir, exist_ok=True)
        stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
        shutil.copy2(p, os.path.join(bdir, f'{os.path.splitext(os.path.basename(p))[0]}-{stamp}.xlsx'))
        for old in sorted(glob.glob(os.path.join(bdir, '*.xlsx')))[:-KEEP_BACKUPS]:
            os.remove(old)
        wb = openpyxl.load_workbook(p)
        ws = wb[str(date.year)] if str(date.year) in wb.sheetnames else new_year_sheet(wb, date.year)
        col = find_column(ws, date) or insert_column(ws, date)
        rows = rows_of(ws)
        fonts = fonts_of(ws, rows)
        for k, v in cells.items():
            set_cell(ws.cell(row=rows[canon(k)], column=col), v, fonts)
        try:
            wb.save(p)
        except PermissionError:
            raise RuntimeError('the spreadsheet is open in another program (Excel?): close it and try again')
        return {'backup': 'schedule/backups'}


class GoogleSheet:
    """The spreadsheet as a Google Sheet, through the small script in tools/google_sheet.gs."""
    def __init__(self):
        import keystore
        self.url, self.secret = keystore.get(SHEET_SECRET, 'url'), keystore.get(SHEET_SECRET, 'secret')
        if not (self.url and self.secret):
            raise RuntimeError('the Google Sheet is not connected yet.')
        self.name = 'the Google Sheet'

    def call(self, **q):
        return call(self.url, self.secret, **q)

    def read(self, date):
        r = self.call(action='read', date=date.isoformat())
        self.name = r.get('title') or self.name
        cells = {canon(k): v for k, v in r['cells'].items()}
        return {'sheet': r['sheetExists'], 'column': r['column'], 'labels': {canon(l): l for l in r['labels']},
                'cells': cells}

    def write(self, date, cells):
        r = self.call(action='write', date=date.isoformat(), cells=cells)
        return {'url': r.get('url')}


def call(url, secret, **q):
    """POST to the sheet's web app. Apps Script answers with a redirect, which urllib follows as a GET."""
    import json, urllib.request, urllib.error
    body = json.dumps(dict(q, secret=secret)).encode()
    req = urllib.request.Request(url, data=body, headers={'Content-Type': 'text/plain;charset=utf-8',
                                                          'User-Agent': 'Mozilla/5.0 (Psalter service planner)'})
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            text = resp.read().decode('utf-8', 'replace')
    except urllib.error.HTTPError as e:
        if e.code in (401, 403):
            raise RuntimeError(f'Google refused the planner (HTTP {e.code}). The web app is probably open only to your '
                               'organization, not to "Anyone": in Apps Script, Deploy › Manage deployments › edit › Who has '
                               'access › Anyone, then Deploy. If "Anyone" isn\'t offered, your Workspace doesn\'t allow it: '
                               'use "Copy for the sheet" and the sheet\'s Psalter menu instead.')
        raise RuntimeError(f'could not reach the Google Sheet ({e})')
    except urllib.error.URLError as e:
        raise RuntimeError(f'could not reach the Google Sheet ({e})')
    try:
        r = json.loads(text)
    except ValueError:
        raise RuntimeError('the Google Sheet answered with a web page, not data: check the web app URL, and that '
                           'it was deployed with access "Anyone"')
    if not r.get('ok'):
        err = r.get('error', 'no')
        raise RuntimeError('the Google Sheet said: ' + err + (' (copy the script again, replace it in the sheet\'s Apps Script and deploy a new version)' if err == 'wrong secret' else ''))
    return r


def backend():
    import lectionary
    v = lectionary.PARISH.get('spreadsheet', '').strip()
    return GoogleSheet() if v.lower() == 'google' else XlsxSheet(path())


def new_secret():
    import secrets
    return secrets.token_urlsafe(24)


def script_text(secret):
    return open(os.path.join(TOOLS, 'google_sheet.gs'), encoding='utf-8').read().replace('__SECRET__', secret)


def connect_google(url, secret):
    """Check the web app answers with this secret, then keep both on this computer."""
    import keystore
    r = call(url.strip(), secret, action='ping')
    keystore.put(SHEET_SECRET, 'url', url.strip())
    keystore.put(SHEET_SECRET, 'secret', secret)
    return r.get('title', 'the sheet')


# ------------------------------------------------------------------ what changes
def plan_changes(sunday_path):
    import build
    sunday = build.parse_sunday(sunday_path)
    if not sunday['date']:
        raise RuntimeError('a psalm without a date does not go in the schedule')
    date = dt.date.fromisoformat(sunday['date'][:10])
    if service.META.get(service.order_of(sunday), {}).get('spreadsheet') is False:
        raise RuntimeError(f"{service.order_of(sunday)} isn't kept in the music schedule (services.txt says so); "
                           'its Sunday file is the record')
    b = backend()
    got = b.read(date)
    labels, current = got['labels'], got['cells']
    notes = []
    new = wanted(sunday, current)
    if not got['sheet']:
        notes.append(f'a new sheet {date.year} will be added')
    elif not got['column']:
        notes.append(f'a new column for {date.isoformat()} will be added in date order')
    order = service.order_of(sunday)
    if not got['column']:
        label = order if order != next(iter(service.ORDERS)) else service.day_label(sunday['date'])
        if label:
            new.setdefault('Proper', label)
    for k in [k for k in new if k not in labels]:
        notes.append(f'no "{k}" row in the sheet: left out ({new.pop(k)})')
    rank = list(labels)
    changes = [(k, current.get(k, ''), v) for k, v in new.items() if (current.get(k, '') or '').strip() != v.strip()]
    changes.sort(key=lambda c: rank.index(c[0]) if c[0] in rank else 999)
    return {'date': date, 'order': order, 'name': b.name, 'sheet': str(date.year), 'changes': changes, 'notes': notes,
            'labels': labels, 'backend': b}


def paste_blob(sunday_path):
    """What the sheet's Psalter menu takes, when the planner can't reach the sheet itself: the cells to write,
    by row name; the sheet works out the column and shows the changes before writing."""
    import build, json
    sunday = build.parse_sunday(sunday_path)
    if not sunday['date']:
        raise RuntimeError('a psalm without a date does not go in the schedule')
    order = service.order_of(sunday)
    if service.META.get(order, {}).get('spreadsheet') is False:
        raise RuntimeError(f"{order} isn't kept in the music schedule (services.txt says so); its Sunday file is the record")
    cells = wanted(sunday, {})
    proper = order if order != next(iter(service.ORDERS)) else service.day_label(sunday['date'])
    return json.dumps({'psalter': 1, 'date': sunday['date'][:10], 'proper': proper, 'cells': cells}, ensure_ascii=False)


def write(sunday_path):
    plan = plan_changes(sunday_path)
    if plan['changes']:
        b = plan['backend']
        res = b.write(plan['date'], {plan['labels'][k]: v for k, _, v in plan['changes']})
        plan.update(res, written=True)
    return plan


def main():
    a = [x for x in sys.argv[1:] if not x.startswith('--')]
    if not a:
        print(__doc__)
        return
    try:
        plan = write(a[0]) if '--write' in sys.argv else plan_changes(a[0])
    except RuntimeError as e:
        raise SystemExit(str(e))
    print(f"{plan['name']}, sheet {plan['sheet']}, {plan['date']} ({plan['order']})")
    for n in plan['notes']:
        print('  note: ' + n)
    for k, old, new in plan['changes']:
        print(f"  {k:<22} {old or '(empty)'}  ->  {new}")
    if not plan['changes']:
        print('  already up to date')
    elif not plan.get('written'):
        print('(nothing written: add --write)')


if __name__ == '__main__':
    main()
